import pandas as pd
import os
import json
from openpyxl.styles import Alignment

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from datetime import datetime
import win32com.client as win32
from dotenv import load_dotenv

# Carrega variáveis de ambiente do arquivo .env (não versionado no Git)
load_dotenv()

# ====================================
# CONFIGURAÇÕES (via .env)
# ====================================

arquivo_entrada = os.getenv("ARQUIVO_ENTRADA")
pasta_saida = os.getenv("PASTA_SAIDA")
tomadores_arquivo = os.getenv("TOMADORES_FILE", "tomadores.json")

_exec_doador_id = os.getenv("EXEC_DOADOR_ID")
_limite_conta_doador = os.getenv("LIMITE_CONTA_DOADOR")

if not arquivo_entrada or not pasta_saida:
    raise RuntimeError(
        "ARQUIVO_ENTRADA e PASTA_SAIDA precisam estar definidos no arquivo .env"
    )

if not _exec_doador_id or not _limite_conta_doador:
    raise RuntimeError(
        "EXEC_DOADOR_ID e LIMITE_CONTA_DOADOR precisam estar definidos no arquivo .env"
    )

EXEC_DOADOR_ID = int(_exec_doador_id)
LIMITE_CONTA_DOADOR = int(_limite_conta_doador)

print("Arquivo de entrada:")
print(os.path.abspath(arquivo_entrada))

# Ler CSV
df = pd.read_csv(
    arquivo_entrada,
    sep=";",
    encoding="latin1"
)

# ====================================
# CADASTRO DE TOMADORES
# ====================================

# Mapa código -> nome da corretora, carregado de arquivo externo
# (não versionado no Git — ver tomadores.example.json)
with open(tomadores_arquivo, "r", encoding="utf-8") as f:
    lista_tomadores = {
        int(codigo): nome
        for codigo, nome in json.load(f).items()
    }

# Identificar colunas
papel = df.columns[0]
quantidade = df.columns[1]
volume = df.columns[2]
taxa = df.columns[3]
exec_doador = df.columns[4]
conta_doador = df.columns[5]
exec_tomador = df.columns[6]

# Converter colunas usadas nos filtros
df[exec_doador] = pd.to_numeric(df[exec_doador], errors="coerce")
df[conta_doador] = pd.to_numeric(df[conta_doador], errors="coerce")

# Converter Quantidade
df[quantidade] = (
    df[quantidade]
    .astype(str)
    .str.replace(".", "", regex=False)
    .str.replace(",", ".", regex=False)
)

df[quantidade] = pd.to_numeric(
    df[quantidade],
    errors="coerce"
)

# Converter Volume
df[volume] = (
    df[volume]
    .astype(str)
    .str.replace(".", "", regex=False)
    .str.replace(",", ".", regex=False)
)

df[volume] = pd.to_numeric(
    df[volume],
    errors="coerce"
)

# Aplicar filtros
df_filtrado = df[
    (df[exec_doador] == EXEC_DOADOR_ID) &
    (df[conta_doador] < LIMITE_CONTA_DOADOR)
]

print(f"\nRegistros filtrados: {len(df_filtrado)}")

# Salvar CSV filtrado
arquivo_csv = os.path.join(
    pasta_saida,
    "NegociosFiltrado.csv"
)

df_filtrado.to_csv(
    arquivo_csv,
    index=False,
    sep=";"
)

print("CSV filtrado criado:")
print(arquivo_csv)

# ====================================
# TABELA DINÂMICA
# ====================================

linhas = []

execucoes_ordenadas = (
    df_filtrado
    .groupby(exec_tomador)[volume]
    .sum()
    .sort_values(ascending=False)
)

for execucao_valor in execucoes_ordenadas.index:

    df_exec = df_filtrado[
        df_filtrado[exec_tomador] == execucao_valor
    ]

    # EXECUÇÃO
    linhas.append([
        str(execucao_valor),
        df_exec[quantidade].sum(),
        df_exec[volume].sum()
    ])

    # PAPÉIS ORDENADOS POR VOLUME
    papeis_ordenados = (
        df_exec
        .groupby(papel)[volume]
        .sum()
        .sort_values(ascending=False)
    )

    for papel_valor in papeis_ordenados.index:

        df_papel = df_exec[
            df_exec[papel] == papel_valor
        ]

        linhas.append([
            f"    {papel_valor}",
            df_papel[quantidade].sum(),
            df_papel[volume].sum()
        ])

        # TAXAS ORDENADAS POR VOLUME
        taxas_ordenadas = (
            df_papel
            .groupby(taxa)[volume]
            .sum()
            .sort_values(ascending=False)
        )

        for taxa_valor in taxas_ordenadas.index:

            df_taxa = df_papel[
                df_papel[taxa] == taxa_valor
            ]

            linhas.append([
                f"        {taxa_valor}",
                df_taxa[quantidade].sum(),
                df_taxa[volume].sum()
            ])

tabela_dinamica = pd.DataFrame(
    linhas,
    columns=[
        "Execução/Papel/Taxa",
        "Quantidade",
        "Volume (R$)"
    ]
)

print(df.columns.tolist())

# ====================================
# RESUMO - TOP 10 EXECUÇÕES
# ====================================

top10_execucao = (
    df_filtrado
    .groupby(exec_tomador, as_index=False)
    .agg({
        quantidade: "sum",
        volume: "sum"
    })
    .sort_values(volume, ascending=False)
)

volume_total_exec = top10_execucao[volume].sum()

top10_execucao["Participação (%)"] = (
    top10_execucao[volume]
    / volume_total_exec
    * 100
).round(2)

top10_execucao["Participação Acumulada (%)"] = (
    top10_execucao["Participação (%)"]
    .cumsum()
    .round(2)
)

top10_execucao.insert(
    0,
    "Ranking",
    range(1, len(top10_execucao) + 1)
)

top10_execucao = top10_execucao.head(10)

# ====================================
# RESUMO - TOP 10 PAPÉIS
# ====================================

top10_papeis = (
    df_filtrado
    .groupby(papel, as_index=False)
    .agg({
        quantidade: "sum",
        volume: "sum"
    })
    .sort_values(volume, ascending=False)
)

volume_total_papeis = top10_papeis[volume].sum()

top10_papeis["Participação (%)"] = (
    top10_papeis[volume]
    / volume_total_papeis
    * 100
).round(2)

top10_papeis["Participação Acumulada (%)"] = (
    top10_papeis["Participação (%)"]
    .cumsum()
    .round(2)
)

top10_papeis.insert(
    0,
    "Ranking",
    range(1, len(top10_papeis) + 1)
)

top10_papeis = top10_papeis.head(10)

# ====================================
# ABA RESUMO
# ====================================

resumo_info = pd.DataFrame({

    "Indicador": [
        "Volume Total (R$)",
        "Total de ações negociadas"
    ],

    "Valor": [
        df_filtrado[volume].sum(),
        df_filtrado[quantidade].sum()
    ]
})

top10_execucao = top10_execucao.rename(
    columns={
        exec_tomador: "Tomador",
        volume: "Volume (R$)"
    }
)

top10_execucao = top10_execucao.drop(
    columns=[quantidade]
)

top10_execucao["Tomador"] = (
    top10_execucao["Tomador"]
    .apply(
        lambda x: f"{x} - {lista_tomadores.get(x, 'Não Cadastrado')}"
    )
)

top10_papeis = top10_papeis.rename(
    columns={
        volume: "Volume (R$)"
    }
)

# ====================================
# EXPORTAÇÃO
# ====================================

arquivo_saida = os.path.join(
    pasta_saida,
    "NegociosDia.xlsx"
)

with pd.ExcelWriter(
    arquivo_saida,
    engine="openpyxl"
) as writer:

    # Dados filtrados
    df_filtrado.to_excel(
        writer,
        sheet_name="Dados Filtrados",
        index=False
    )

    # Resumo
    resumo_info.to_excel(
        writer,
        sheet_name="Resumo",
        startrow=0,
        index=False
    )

    top10_execucao.to_excel(
        writer,
        sheet_name="Resumo",
        startrow=8,
        index=False
    )

    top10_papeis.to_excel(
        writer,
        sheet_name="Resumo",
        startrow=25,
        index=False
    )

    # Estrutura igual ao Resumo Diário
    tabela_dinamica.to_excel(
        writer,
        sheet_name="Tabela Dinamica",
        index=False
    )

print("\nArquivo criado com sucesso:")
print(arquivo_saida)

# ====================================
# FORMATAÇÃO DO EXCEL
# ====================================

wb = load_workbook(arquivo_saida)

# Cores
azul_petroleo = PatternFill(
    start_color="0C4F57",
    end_color="0C4F57",
    fill_type="solid"
)

cinza = PatternFill(
    start_color="D9D9D9",
    end_color="D9D9D9",
    fill_type="solid"
)

fonte_branca = Font(
    color="FFFFFF",
    bold=True
)

fonte_negrito = Font(
    bold=True
)

# ====================================
# FORMATAR ABAS DADOS FILTRADOS E TABELA DINAMICA
# ====================================

# Congelar painéis
for ws in wb.worksheets:

    if ws.title != "Resumo":
        ws.freeze_panes = "A2"

for ws in wb.worksheets:

    # Cabeçalho
    if ws.title == "Resumo":
        for cell in ws[1][:2]:
            cell.fill = azul_petroleo
            cell.font = fonte_branca

    else:
        for cell in ws[1]:
            cell.fill = azul_petroleo
            cell.font = fonte_branca

    # Filtro
    if ws.title != "Resumo":
        ws.auto_filter.ref = ws.dimensions

    # Ajuste automático das colunas
    for coluna in ws.columns:

        max_length = 0

        letra = get_column_letter(
            coluna[0].column
        )

        for cell in coluna:

            try:
                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )
            except:
                pass

        ws.column_dimensions[letra].width = (
            max_length + 3
        )

# ====================================
# FORMATAR TABELA DINÂMICA
# ====================================

ws = wb["Tabela Dinamica"]

for row in ws.iter_rows(min_row=2):

    valor = row[0].value

    if valor is None:
        continue

    valor = str(valor)

    # EXECUÇÃO
    if not valor.startswith(" "):

        for cell in row:
            cell.fill = azul_petroleo
            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )

    # PAPEL
    elif valor.startswith("    ") and not valor.startswith("        "):

        for cell in row:
            cell.fill = cinza
            cell.font = fonte_negrito

# ====================================
# FORMATAÇÃO NUMÉRICA
# ====================================

for ws in wb.worksheets:

    for row in ws.iter_rows(min_row=2):

        for cell in row:

            if isinstance(cell.value, (int, float)):

                if cell.column == 3:
                    cell.number_format = '#,##0'

                if cell.column == 4:
                    cell.number_format = '#,##0.00'

                # Quantidade e Volume
                if ws.title == "Tabela Dinamica":
                    if cell.column == 2:
                        cell.number_format = '#,##0'
                    if cell.column == 3:
                        cell.number_format = '#,##0.00'

# ====================================
# FORMATAR ABA RESUMO
# ====================================

ws = wb["Resumo"]

# ------------------------------------
# Indicadores
# ------------------------------------

for linha in range(2, 4):

    ws[f"A{linha}"].font = Font(
        bold=True
    )

    ws[f"B{linha}"].font = Font(
        bold=True,
        size=12
    )

# Formatação dos indicadores
ws["B2"].number_format = '#,##0.00'
ws["B3"].number_format = '#,##0'

# ------------------------------------
# Título Top 10 Execuções
# ------------------------------------

ws.merge_cells("A8:F8")

titulo = ws["A8"]

titulo.value = "TOP 10 EXECUÇÕES POR VOLUME"
titulo.fill = azul_petroleo
titulo.font = fonte_branca
titulo.alignment = Alignment(horizontal="center")

# Cabeçalhos da tabela (linha 9)
for coluna in ["A", "B", "C", "D", "E", "F"]:

    ws[f"{coluna}9"].fill = azul_petroleo
    ws[f"{coluna}9"].font = fonte_branca

# ------------------------------------
# Título Top 10 Papéis
# ------------------------------------

ws.merge_cells("A25:F25")

titulo = ws["A25"]

titulo.value = "TOP 10 PAPÉIS POR VOLUME"
titulo.fill = azul_petroleo
titulo.font = fonte_branca
titulo.alignment = Alignment(horizontal="center")

# Cabeçalhos da tabela (linha 26)
for coluna in ["A", "B", "C", "D", "E", "F"]:

    ws[f"{coluna}26"].fill = azul_petroleo
    ws[f"{coluna}26"].font = fonte_branca

# Reordenar abas
aba_resumo = wb["Resumo"]
aba_tabela = wb["Tabela Dinamica"]
aba_dados = wb["Dados Filtrados"]

wb._sheets = [
    aba_resumo,
    aba_tabela,
    aba_dados
]

# Ocultar aba de dados
wb["Dados Filtrados"].sheet_state = "hidden"

wb.save(arquivo_saida)

print("Formatação aplicada com sucesso.")

# ====================================
# ENVIO DE E-MAIL OUTLOOK
# ====================================

def formato_br(valor):
    return (
        f"{valor:,.2f}"
        .replace(",", "X")
        .replace(".", ",")
        .replace("X", ".")
    )

def formato_br_int(valor):
    return (
        f"{valor:,.0f}"
        .replace(",", ".")
    )

# ====================================
# PREPARAR TABELA TOP 10 TOMADORES
# ====================================

top10_execucao_email = top10_execucao.copy()

top10_execucao_email["Volume (R$)"] = (
    top10_execucao_email["Volume (R$)"]
    .apply(formato_br)
)

html_execucao = top10_execucao_email.to_html(
    index=False,
    border=1,
    justify="center"
)

# ====================================
# PREPARAR TABELA TOP 10 PAPÉIS
# ====================================

top10_papeis_email = top10_papeis.copy()

top10_papeis_email["Volume (R$)"] = (
    top10_papeis_email["Volume (R$)"]
    .apply(formato_br)
)

top10_papeis_email[quantidade] = (
    top10_papeis_email[quantidade]
    .apply(formato_br_int)
)

html_papeis = top10_papeis_email.to_html(
    index=False,
    border=1,
    justify="center"
)

# ====================================
# INDICADORES
# ====================================

volume_total_html = formato_br(
    df_filtrado[volume].sum()
)

quantidade_total_html = formato_br_int(
    df_filtrado[quantidade].sum()
)

# ====================================
# DATA DO ASSUNTO
# ====================================

data_hoje = datetime.now().strftime("%d/%m/%Y")

# ====================================
# HTML DO E-MAIL
# ====================================

html_body = f"""
<html>

<body style="font-family:Calibri;font-size:11pt;">

<p>Boa noite,</p>

<p>Segue Resumo Diário de BTC.</p>

<h3>Resumo</h3>

<table border="1" cellpadding="5" cellspacing="0">
<tr>
    <th>Indicador</th>
    <th>Valor</th>
</tr>

<tr>
    <td>Volume Total (R$)</td>
    <td>{volume_total_html}</td>
</tr>

<tr>
    <td>Total de ações negociadas</td>
    <td>{quantidade_total_html}</td>
</tr>

</table>

<br>

<h3>Top 10 Tomadores por Volume</h3>

{html_execucao}

<br>

<h3>Top 10 Papéis por Volume</h3>

{html_papeis}

<br>

<p>
O relatório completo encontra-se em anexo.
</p>

</body>
</html>
"""

# ====================================
# OUTLOOK
# ====================================

outlook = win32.Dispatch("Outlook.Application")

mail = outlook.CreateItem(0)

destinatarios = os.getenv("EMAIL_DESTINATARIOS")
copia = os.getenv("EMAIL_COPIA")

if not destinatarios:
    raise RuntimeError(
        "EMAIL_DESTINATARIOS precisa estar definido no arquivo .env"
    )

mail.To = destinatarios
mail.CC = copia

mail.Subject = (
    f"Resumo Diário BTC - {data_hoje}"
)

mail.HTMLBody = html_body

mail.Attachments.Add(arquivo_saida)

# Abre para conferência
mail.Display()

# Para envio automático usar:
# mail.Send()

print("E-mail criado com sucesso.")
